"""Block 4: the parameterised scoring engine, checked against real results.

    python3 tests/test_scoring.py

Costs nothing. No Anthropic call is made anywhere in this file: `score_job` is
the only function that talks to the API and it is never invoked. Sub-scores and
flags come from the existing .xlsx results; JobPosting objects for prompt
building come from the saved connector dumps.

Two differential tests, deliberately separate:

  1. LOGIC      - replaying the adjustments and thresholds against the base
                  score the pipeline itself recorded. This is the behaviour
                  being ported, and it must match exactly.
  2. ARITHMETIC - the weighted average, which used to be done by the model and
                  now happens in Python. Here the two are EXPECTED to differ,
                  because the model's arithmetic was frequently wrong. What is
                  asserted is that ours is mathematically exact.
"""

from __future__ import annotations

import re
import sys
from decimal import Decimal
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

from openpyxl import load_workbook  # noqa: E402

from jobscout.connectors import Secrets, build_connectors  # noqa: E402
from jobscout.jobs import JobPosting  # noqa: E402
from jobscout.profile import load_profile  # noqa: E402
from jobscout.prompts import build_eval_prompt, build_eval_system  # noqa: E402
from jobscout.scoring import (  # noqa: E402
    Adjustment,
    ResponseError,
    apply_adjustments,
    compute_verdict,
    parse_response,
    round1,
    weighted_base,
)

CHECKS = 0
FAILURES: list[str] = []


def check(label: str, actual, expected) -> None:
    global CHECKS
    CHECKS += 1
    if actual == expected:
        print(f"  OK    {label}")
        return
    FAILURES.append(label)
    print(f"  FAIL  {label}")
    print(f"          expected: {expected!r}")
    print(f"          actual  : {str(actual)[:300]!r}")


def section(title: str) -> None:
    print(f"\n{'=' * 78}\n{title}\n{'=' * 78}")


gabriel = load_profile(REPO / "profiles" / "gabriel.json", base_dir=REPO)
camila = load_profile(REPO / "profiles" / "camila.json", base_dir=REPO)
G_CV = gabriel.load_cv(REPO)
C_CV = camila.load_cv(REPO)
BCN = JobPosting(title="t", location="Barcelona")


# ---------------------------------------------------------------------------
# Read the real results back out of the .xlsx files
# ---------------------------------------------------------------------------

def read_rows(path: Path, breakdown_col: str, score_col: str, verdict_col: str,
              redflag_col: str, signal_col: str) -> list[dict]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["Job Scout"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(header, raw))
        breakdown = str(row.get(breakdown_col) or "")
        pairs = re.findall(r"(\w+):(\d+)", breakdown)     # "Tech:2 Loc:6 ..."
        total = re.search(r"→\s*([\d.]+)", breakdown)     # "... → 4.1"
        if not pairs or not total:
            continue
        try:
            score = float(row.get(score_col))
        except (TypeError, ValueError):
            continue
        rows.append({
            "labels": {k: int(v) for k, v in pairs},
            "reported_base": float(total.group(1)),
            "final_score": score,
            "verdict": str(row.get(verdict_col) or ""),
            "markers": f"{row.get(redflag_col) or ''} {row.get(signal_col) or ''}",
        })
    wb.close()
    return rows


g_rows = read_rows(REPO / "job_scout_results.xlsx",
                   "Score breakdown", "Score", "Verdict", "Red flags", "Match signals")
c_rows = read_rows(REPO / "job_scout_camila" / "job_scout_camila_results.xlsx",
                   "Desglose puntuación", "Puntuación", "Veredicto",
                   "Banderas rojas", "Señales positivas")

G_LABEL_TO_KEY = {d.label: d.key for d in gabriel.dimensions}
C_LABEL_TO_KEY = {d.label: d.key for d in camila.dimensions}


def sub_scores_for(labels: dict[str, int], mapping: dict[str, str]) -> dict[str, int]:
    return {mapping[k]: v for k, v in labels.items() if k in mapping}


def declared_adjustments(profile) -> list[tuple]:
    return (
        [(f.key, f.adjustment, f.note, f.target, "llm_flag")
         for f in profile.llm_flags]
        + [(p.key, p.adjustment, p.note,
            "match_signals" if p.adjustment > 0 else "red_flags", "keyword_penalty")
           for p in profile.keyword_penalties]
    )


def fired_adjustments(row: dict, profile) -> list[Adjustment]:
    """Which adjustments the original pipeline applied to this row.

    Matched on the exact marker the pipeline wrote ("[penalty -3.0] <note>"),
    not on the note text alone: the model also writes free-form red flags that
    can contain the same words, which would over-detect.
    """
    out = []
    for key, amount, note, target, origin in declared_adjustments(profile):
        marker = (f"[penalty −{abs(amount):.1f}] {note}" if amount < 0
                  else f"[+{amount:.1f}] {note}")
        if marker in row["markers"]:
            out.append(Adjustment(key=key, amount=amount, note=note,
                                  target=target, origin=origin))
    return out


VERDICT_BACK = {
    "OUI": "YES", "OPPORTUNISTE": "MAYBE", "NON": "NO",
    "SÍ": "YES", "OPORTUNISTA": "MAYBE", "NO": "NO",
}


# ===========================================================================
section("DATA — rows recovered from the existing result files")

check("gabriel rows parsed", len(g_rows) > 900, True)
check("camila rows parsed", len(c_rows) > 50, True)
print(f"\n  gabriel: {len(g_rows)} rows | camila: {len(c_rows)} rows")
check("gabriel breakdown labels map to dimensions",
      sorted({k for r in g_rows for k in r["labels"]}), sorted(G_LABEL_TO_KEY))
check("camila breakdown labels map to dimensions",
      sorted({k for r in c_rows for k in r["labels"]}), sorted(C_LABEL_TO_KEY))


# ===========================================================================
section("DIFFERENTIAL 1 — adjustments and thresholds vs the pipeline")

for label, rows, profile in (("gabriel", g_rows, gabriel),
                             ("camila", c_rows, camila)):
    score_diff, verdict_diff, self_inconsistent, with_adj = [], [], [], 0
    for row in rows:
        adjustments = fired_adjustments(row, profile)
        with_adj += bool(adjustments)
        score = apply_adjustments(row["reported_base"], adjustments)
        if score != row["final_score"]:
            score_diff.append((row["reported_base"],
                               [a.label for a in adjustments],
                               score, row["final_score"]))

        t = profile.thresholds
        verdict = ("YES" if score > t.yes_above
                   else "MAYBE" if score >= t.maybe_above else "NO")
        recorded = VERDICT_BACK.get(row["verdict"], row["verdict"])

        # The old pipeline let the model decide the verdict, so a row can carry
        # a verdict that contradicts the score printed next to it. Those rows
        # say nothing about our logic - count them, but compare only the rest.
        consistent = ("YES" if row["final_score"] > t.yes_above
                      else "MAYBE" if row["final_score"] >= t.maybe_above else "NO")
        if recorded != consistent:
            self_inconsistent.append((row["final_score"], recorded, consistent))
        elif verdict != recorded:
            verdict_diff.append((score, verdict, recorded))

    print(f"\n  {label}: {len(rows)} rows replayed, {with_adj} carried adjustments")

    # The only score divergences allowed are exact .5 rounding boundaries,
    # where the pipeline used Python's round() (binary, banker's) and we use
    # round-half-up. Anything else is a real failure.
    for base, labels, mine, theirs in score_diff[:3]:
        print(f"    base {base} {labels} -> ours {mine} / pipeline {theirs}"
              f"   (differs by {abs(round(mine - theirs, 2))})")
    check(f"{label}: every score divergence is a 0.1 rounding boundary",
          [d for d in score_diff if abs(round(d[2] - d[3], 2)) != 0.1], [])
    check(f"{label}: at most one such boundary row", len(score_diff) <= 1, True)

    print(f"    rows whose recorded verdict contradicts their own recorded "
          f"score: {len(self_inconsistent)}")
    for sc, rec, cons in self_inconsistent[:3]:
        print(f"      score {sc} -> pipeline said {rec}, its own thresholds give {cons}")
    check(f"{label}: verdict reproduced on every self-consistent row",
          verdict_diff, [])

# A worked example, so the mechanism is visible and not just asserted.
worked = next(r for r in c_rows if len(fired_adjustments(r, camila)) >= 3)
adjs = fired_adjustments(worked, camila)
print(f"\n  worked example: base {worked['reported_base']}")
for a in adjs:
    print(f"    {a.label}")
print(f"    -> {apply_adjustments(worked['reported_base'], adjs)} "
      f"(pipeline recorded {worked['final_score']})")


# ===========================================================================
section("DIFFERENTIAL 2 — the arithmetic the model used to do itself")

for label, rows, profile, mapping in (("gabriel", g_rows, gabriel, G_LABEL_TO_KEY),
                                      ("camila", c_rows, camila, C_LABEL_TO_KEY)):
    wrong, deltas, checked, inexact = 0, [], 0, []
    for row in rows:
        sub = sub_scores_for(row["labels"], mapping)
        if len(sub) != len(profile.dimensions):
            continue
        checked += 1
        ours = weighted_base(profile, sub)
        # Independent check, computed a different way with exact decimals.
        exact = sum(Decimal(str(sub[d.key])) * Decimal(str(d.weight))
                    for d in profile.dimensions)
        if abs(Decimal(str(ours)) - exact) > Decimal("0.05"):
            inexact.append((sub, ours, float(exact)))
        if ours != row["reported_base"]:
            wrong += 1
            deltas.append(round(ours - row["reported_base"], 1))

    pct = 100.0 * wrong / max(checked, 1)
    print(f"\n  {label}: {checked} rows")
    print(f"    our weighted average is mathematically exact on all {checked}")
    print(f"    the model's own total disagreed on {wrong} of them ({pct:.0f}%)")
    if deltas:
        under = sum(1 for d in deltas if d > 0)
        print(f"    model error: min {min(deltas):+.1f}  max {max(deltas):+.1f}  "
              f"mean {sum(deltas) / len(deltas):+.2f}")
        print(f"    the model scored LOW in {under}/{len(deltas)} "
              f"({100.0 * under / len(deltas):.0f}%) of the disagreements")
    check(f"{label}: our arithmetic is exact on every row", inexact, [])


# ===========================================================================
section("ARITHMETIC — weights, clamping, rounding")

# 9*0.40 + 10*0.25 + 8*0.20 + 10*0.15 = 9.2. The notebook printed 9.4 for
# these sub-scores, because the model did the arithmetic. 9.2 is correct.
check("gabriel weighted average",
      weighted_base(gabriel, {"tech_fit": 9, "location": 10,
                              "company_size": 8, "seniority": 10}), 9.2)
check("camila weighted average",
      weighted_base(camila, {"perfil_fit": 8, "location": 7,
                             "entidad_fit": 9, "condiciones": 7}), 7.9)
check("round1 rounds half away from zero", round1(7.85), 7.9)
check("built-in round() would have given 7.8", round(7.85, 1), 7.8)

missing = False
try:
    weighted_base(gabriel, {"tech_fit": 5})
except ValueError:
    missing = True
check("a missing sub-score raises rather than defaulting", missing, True)

check("adjustments are summed, then clamped once at 0",
      apply_adjustments(7.9, [Adjustment(key="a", amount=-2.5, note="x"),
                              Adjustment(key="b", amount=1.0, note="y"),
                              Adjustment(key="c", amount=-3.0, note="z"),
                              Adjustment(key="d", amount=-3.0, note="w")]), 0.4)
check("summing is order-independent",
      apply_adjustments(7.9, [Adjustment(key="b", amount=1.0, note="y"),
                              Adjustment(key="d", amount=-3.0, note="w"),
                              Adjustment(key="a", amount=-2.5, note="x"),
                              Adjustment(key="c", amount=-3.0, note="z")]), 0.4)
check("the floor is 0, never negative",
      apply_adjustments(1.0, [Adjustment(key="a", amount=-9.0, note="x")]), 0.0)
check("the ceiling is 10",
      apply_adjustments(9.8, [Adjustment(key="a", amount=1.0, note="x")]), 10.0)

check("a keyword penalty fires from the job, not from the model",
      [a.key for a in compute_verdict(
          camila, {"perfil_fit": 9, "location": 9, "entidad_fit": 9,
                   "condiciones": 9}, {},
          JobPosting(title="t", location="Mataró")).adjustments],
      ["corredor_maresme_r1"])
check("and not when the location is fine",
      compute_verdict(camila, {"perfil_fit": 9, "location": 9, "entidad_fit": 9,
                               "condiciones": 9}, {}, BCN).adjustments, [])


# ===========================================================================
section("THRESHOLDS — one source of truth")

check("exactly 7.0 is MAYBE (strictly greater is required for YES)",
      compute_verdict(gabriel, {"tech_fit": 7, "location": 7, "company_size": 7,
                                "seniority": 7}, {}, BCN).verdict, "MAYBE")
check("exactly 5.0 is MAYBE (greater or equal)",
      compute_verdict(gabriel, {"tech_fit": 5, "location": 5, "company_size": 5,
                                "seniority": 5}, {}, BCN).verdict, "MAYBE")
check("4.9 is NO",
      compute_verdict(gabriel, {"tech_fit": 5, "location": 5, "company_size": 5,
                                "seniority": 4}, {}, BCN).verdict, "NO")

strict = gabriel.model_copy(update={
    "thresholds": gabriel.thresholds.model_copy(update={"yes_above": 9.0})})
eight = {"tech_fit": 8, "location": 8, "company_size": 8, "seniority": 8}
check("changing the profile's threshold changes the verdict",
      compute_verdict(strict, eight, {}, BCN).verdict, "MAYBE")
check("the same sub-scores are YES under the real profile",
      compute_verdict(gabriel, eight, {}, BCN).verdict, "YES")
check("and the score itself is unchanged by the threshold",
      compute_verdict(strict, eight, {}, BCN).score,
      compute_verdict(gabriel, eight, {}, BCN).score)


# ===========================================================================
section("PROMPT — no weights, no thresholds, no verdict vocabulary")

sample = JobPosting(title="Educador/a social", company="Fundació Test",
                    location="Barcelona",
                    summary="Contrato indefinido. Trabajo comunitario. " * 8)
g_prompt = build_eval_prompt(gabriel, sample, G_CV)
c_prompt = build_eval_prompt(camila, sample, C_CV)

for label, prompt, profile, cv in (("gabriel", g_prompt, gabriel, G_CV),
                                   ("camila", c_prompt, camila, C_CV)):
    full = prompt + "\n" + build_eval_system(profile)

    check(f"{label}: no verdict threshold value appears",
          [t for t in (str(profile.thresholds.yes_above),
                       str(profile.thresholds.maybe_above)) if t in full], [])
    check(f"{label}: no weight value appears",
          [str(d.weight) for d in profile.dimensions if str(d.weight) in full], [])
    check(f"{label}: no weighting formula",
          [t for t in ("weighted_score", "weighted total", "*0.", "* 0.")
           if t in full], [])
    check(f"{label}: no verdict vocabulary",
          [t for t in ("OUI", "OPPORTUNISTE", "NON", "OPORTUNISTA", "verdict")
           if t in full], [])

    check(f"{label}: every dimension key is present",
          all(d.key in prompt for d in profile.dimensions), True)
    check(f"{label}: every rubric is present",
          all(d.rubric[:60] in prompt for d in profile.dimensions), True)
    check(f"{label}: every flag definition is present",
          all(f.definition[:50] in prompt for f in profile.llm_flags), True)
    check(f"{label}: the CV is included", cv[:40] in prompt, True)
    check(f"{label}: sub_scores is requested", '"sub_scores"' in prompt, True)

check("camila's prompt asks for all 5 flags",
      all(f.key in c_prompt for f in camila.llm_flags), True)
check("gabriel has no flags block (he declares none)", "## Flags" in g_prompt, False)
check("the Maresme penalty is never shown to the model",
      "Maresme" in c_prompt or "Mataró" in c_prompt, False)
check("extra output fields are requested",
      all(f in g_prompt for f in gabriel.extra_output_fields), True)
check("the dead cover_letter boolean is gone", '"cover_letter"' in g_prompt, False)
print(f"\n  gabriel prompt {len(g_prompt)} chars | camila prompt {len(c_prompt)} chars")


# ===========================================================================
section("PROMPT — built from real jobs in the saved dumps (no API call)")

conns_c = build_connectors(camila, REPO / "dumps")
ij_jobs = conns_c["infojobs_apify"].fetch(
    camila.source("infojobs_apify").model_copy(update={"reuse_dump": True}),
    Secrets())[:25]
lengths = [len(build_eval_prompt(camila, j, C_CV)) for j in ij_jobs]
check("prompts build for 25 real InfoJobs postings", len(lengths), 25)
check("each contains its own job title",
      all(j.title in build_eval_prompt(camila, j, C_CV) for j in ij_jobs), True)
check("none is suspiciously short", min(lengths) > 2000, True)
print(f"  InfoJobs prompt length min/max: {min(lengths)} / {max(lengths)}")

conns_g = build_connectors(gabriel, REPO / "dumps")
li_jobs = conns_g["linkedin_apify"].fetch(
    gabriel.source("linkedin_apify").model_copy(update={"reuse_dump": True}),
    Secrets())[:25]
li_lengths = [len(build_eval_prompt(gabriel, j, G_CV)) for j in li_jobs]
check("prompts build for 25 real LinkedIn postings", len(li_lengths), 25)
print(f"  LinkedIn prompt length min/max: {min(li_lengths)} / {max(li_lengths)}")


# ===========================================================================
section("RESPONSE PARSING — including the retry path")

good = ('{"sub_scores": {"tech_fit": 8, "location": 9, "company_size": 6, '
        '"seniority": 7}, "one_liner": "Good fit.", "match_signals": ["Python"], '
        '"gaps": [], "red_flags": [], "salary_range_market": "60k"}')
sub, flags, payload = parse_response(good, gabriel)
check("parses a clean reply", sub,
      {"tech_fit": 8, "location": 9, "company_size": 6, "seniority": 7})
check("no flags for a profile that declares none", flags, {})
check("prose is preserved", payload["one_liner"], "Good fit.")
check("strips a ```json fence",
      parse_response("```json\n" + good + "\n```", gabriel)[0]["tech_fit"], 8)
check("strips a bare ``` fence",
      parse_response("```\n" + good + "\n```", gabriel)[0]["tech_fit"], 8)
check("coerces a float sub-score",
      parse_response(good.replace('"tech_fit": 8', '"tech_fit": 8.4'),
                     gabriel)[0]["tech_fit"], 8)
check("clamps an out-of-range sub-score",
      parse_response(good.replace('"tech_fit": 8', '"tech_fit": 47'),
                     gabriel)[0]["tech_fit"], 10)

for label, bad in (
    ("truncated JSON", good[:60]),
    ("a JSON array instead of an object", "[1, 2, 3]"),
    ("a missing dimension", '{"sub_scores": {"tech_fit": 8}}'),
    ("a non-numeric sub-score", good.replace('"tech_fit": 8', '"tech_fit": "high"')),
):
    raised = False
    try:
        parse_response(bad, gabriel)
    except ResponseError:
        raised = True
    check(f"rejects {label}", raised, True)

flagged = ('{"sub_scores": {"perfil_fit": 8, "location": 7, "entidad_fit": 9, '
           '"condiciones": 7}, "flags": {"catalan_imprescindible": true}, '
           '"one_liner": "x", "tipo_contrato": "indefinido"}')
sub_c, flags_c, payload_c = parse_response(flagged, camila)
check("reads the flags the model answered", flags_c["catalan_imprescindible"], True)
check("defaults unanswered flags to false", flags_c["horario_atipico"], False)
check("all five flags are always present", len(flags_c), 5)
check("extra output fields survive", payload_c["tipo_contrato"], "indefinido")

result = compute_verdict(camila, sub_c, flags_c, BCN)
check("end to end: base 7.9 minus 3.0 catalan", result.score, 4.9)
check("and lands on NO", result.verdict, "NO")
check("the adjustment is recorded for display",
      [a.label for a in result.adjustments], ["[-3.0] català imprescindible"])


# ===========================================================================
print()
print("=" * 78)
if FAILURES:
    print(f"FAILED — {len(FAILURES)} of {CHECKS} checks:")
    for f in FAILURES:
        print(f"  - {f}")
    print("=" * 78)
    sys.exit(1)
print(f"PASSED — {CHECKS} checks. No Anthropic API call was made.")
print("=" * 78)
