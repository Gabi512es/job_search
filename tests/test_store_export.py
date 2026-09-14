"""Block 5: persistence, the job_results mapping, and the secondary Excel export.

    python3 tests/test_store_export.py

Costs nothing: no Anthropic call, no Apify call, no network. Results are
rebuilt from the rows already in the two .xlsx files.

The non-regression target for the export is the notebooks' own COLUMNS and
TRACKER_COLS_DEF lists, read live out of the notebooks - not the .xlsx files on
disk, which were written by older versions with fewer columns.
"""

from __future__ import annotations

import ast
import shutil
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

from openpyxl import load_workbook  # noqa: E402

from jobscout.export.computed import REGISTRY, priority_key, resolve  # noqa: E402
from jobscout.export.excel import export_xlsx  # noqa: E402
from jobscout.jobs import JobPosting  # noqa: E402
from jobscout.profile import load_profile  # noqa: E402
from jobscout.scoring import ScoredJob  # noqa: E402
from jobscout.store.base import Application, JobResult, RunRecord, to_job_result  # noqa: E402
from jobscout.store.json_store import JsonStore  # noqa: E402
from jobscout.store.supabase_store import SupabaseStore  # noqa: E402
from tests.legacy_values import CAMILA_NB, GABRIEL_NB, cell_sources  # noqa: E402

CHECKS = 0
FAILURES: list[str] = []
TMP = Path("/private/tmp/claude-501/-Users-gabrielernoult-Desktop-GIT-Repos-Job-Search"
           "/a697afae-6c9e-4187-870d-45749f73a1a8/scratchpad/block5")


def check(label: str, actual, expected) -> None:
    global CHECKS
    CHECKS += 1
    if actual == expected:
        print(f"  OK    {label}")
        return
    FAILURES.append(label)
    print(f"  FAIL  {label}")
    print(f"          expected: {str(expected)[:220]!r}")
    print(f"          actual  : {str(actual)[:220]!r}")


def section(title: str) -> None:
    print(f"\n{'=' * 78}\n{title}\n{'=' * 78}")


shutil.rmtree(TMP, ignore_errors=True)
TMP.mkdir(parents=True, exist_ok=True)

gabriel = load_profile(REPO / "profiles" / "gabriel.json", base_dir=REPO)
camila = load_profile(REPO / "profiles" / "camila.json", base_dir=REPO)


# ---------------------------------------------------------------------------
# The notebooks' own column definitions, read live
# ---------------------------------------------------------------------------

def notebook_columns(nb_path: Path, cell_index: int, name: str) -> list[tuple[str, int]]:
    """Extract COLUMNS / TRACKER_COLS_DEF as [(label, width), ...]."""
    tree = ast.parse(cell_sources(nb_path)[cell_index])
    for node in ast.walk(tree):
        if isinstance(node, ast.Assign):
            for target in node.targets:
                if isinstance(target, ast.Name) and target.id == name:
                    return [tuple(ast.literal_eval(e)) for e in node.value.elts]
    raise AssertionError(f"{name} not found")


G_COLS = notebook_columns(GABRIEL_NB, 15, "COLUMNS")
G_TRACKER = notebook_columns(GABRIEL_NB, 15, "TRACKER_COLS_DEF")
C_COLS = notebook_columns(CAMILA_NB, 15, "COLUMNS")
C_TRACKER = notebook_columns(CAMILA_NB, 15, "TRACKER_COLS_DEF")


# ===========================================================================
section("STORE — JsonStore behaviour")

store = JsonStore(TMP / "store")

r1 = JobResult(user_id="u1", run_id="run1", job_key="k1", title="A",
               url="https://x/1", verdict="YES", score=8.0)
r2 = JobResult(user_id="u1", run_id="run1", job_key="k2", title="B",
               url="https://x/2", verdict="NO", score=3.0)

check("inserts new results", store.save_results("u1", [r1, r2]), 2)
check("the unique (user_id, url) constraint holds",
      store.save_results("u1", [r1]), 0)
check("results come back", len(store.get_results("u1")), 2)
check("filtered by verdict",
      [r.title for r in store.get_results("u1", verdicts=["YES"])], ["A"])
check("known_urls reflects what is stored",
      store.known_urls("u1"), {"https://x/1", "https://x/2"})

store.mark_seen("u1", {"k1", "k2"})
store.mark_seen("u1", {"k2", "k3"})
check("seen keys accumulate without duplicating",
      store.seen_keys("u1"), {"k1", "k2", "k3"})

# Per-user isolation: the shape RLS will enforce in Supabase.
r3 = JobResult(user_id="u2", run_id="run1", job_key="k1", title="A",
               url="https://x/1", verdict="YES", score=8.0)
check("a second user can hold the same URL", store.save_results("u2", [r3]), 1)
check("and sees only their own", len(store.get_results("u2")), 1)
check("without affecting the first", len(store.get_results("u1")), 2)
check("seen keys are per user", store.seen_keys("u2"), set())

mismatched = False
try:
    store.save_results("u1", [r3])
except ValueError:
    mismatched = True
check("refuses a result belonging to another user", mismatched, True)

for bad in ("../escape", "a/b", ""):
    rejected = False
    try:
        store.seen_keys(bad)
    except ValueError:
        rejected = True
    check(f"rejects unsafe user_id {bad!r}", rejected, True)

run = RunRecord(run_id="run1", user_id="u1", profile_id="gabriel",
                status="OK", cost_decision="OK", cost_fingerprint="abc123",
                counts={"to_score": 12})
store.save_run(run)
check("a run round-trips", store.get_run("u1", "run1").cost_fingerprint, "abc123")
check("an unknown run is None", store.get_run("u1", "nope"), None)
store.save_run(run.model_copy(update={"status": "finished"}))
check("re-saving a run replaces it, not duplicates",
      len(store._read("u1", "runs.json", [])), 1)

section("STORE — applications survive runs")

store.save_application(Application(user_id="u1", job_url="https://x/1",
                                   status="Entretien", notes="RDV mardi"))
check("an application is stored",
      store.get_applications("u1")["https://x/1"].status, "Entretien")
check("updated_at is filled in",
      bool(store.get_applications("u1")["https://x/1"].updated_at), True)

# The bug this design fixes: in the notebooks the Tracker sheet was deleted and
# rebuilt every run, so hand-edited statuses were lost.
store.save_results("u1", [JobResult(user_id="u1", run_id="run2", job_key="k9",
                                    title="C", url="https://x/9", verdict="YES")])
check("a later run does not touch the application",
      store.get_applications("u1")["https://x/1"].status, "Entretien")
check("nor its notes",
      store.get_applications("u1")["https://x/1"].notes, "RDV mardi")

section("STORE — JsonStore and SupabaseStore are interchangeable")

# SupabaseStore was a shell at block 5 and is implemented now; its own
# behaviour is covered by tests/test_supabase_store.py. What matters here is
# that the two remain substitutable, and that the Supabase one cannot be
# constructed by accident without credentials.
import inspect  # noqa: E402

json_methods = sorted(m for m in dir(JsonStore)
                      if not m.startswith("_") and callable(getattr(JsonStore, m)))
check("SupabaseStore implements every JsonStore method",
      [m for m in json_methods if not hasattr(SupabaseStore, m)], [])
check("with identical signatures",
      [m for m in json_methods
       if list(inspect.signature(getattr(JsonStore, m)).parameters)
       != list(inspect.signature(getattr(SupabaseStore, m)).parameters)], [])

raised = False
try:
    SupabaseStore("", "")
except Exception as exc:
    raised = "SUPABASE_URL" in str(exc)
check("SupabaseStore refuses to build without credentials", raised, True)


# ===========================================================================
section("MAPPING — ScoredJob -> job_results row")

job = JobPosting(title="Educador/a social", company="Fundació Test",
                 url="https://example.org/42?utm_source=x", source="xarxanet",
                 location="Barcelona", published="2026-09-01")
scored = ScoredJob(
    title=job.title, company=job.company, url=job.url, source=job.source,
    location=job.location, published=job.published,
    base_score=7.9, score=4.9, verdict="NO",
    sub_scores={"perfil_fit": 8, "location": 7, "entidad_fit": 9, "condiciones": 7},
    flags={"catalan_imprescindible": True, "horario_atipico": False},
    one_liner="Buen encaje pero catalán requerido.",
    match_signals=["trabajo comunitario"], gaps=["catalán"],
    red_flags=["[-3.0] català imprescindible"],
    salary_range_market="€22k-26k", extra={"tipo_contrato": "indefinido"},
    evaluated_at="2026-09-12T10:00:00",
)
row = to_job_result(scored, job, user_id="u1", run_id="run7")

check("job_key comes from the posting, not the score", row.job_key, job.key)
check("verdict carried over", row.verdict, "NO")
check("score carried over", row.score, 4.9)
check("base score kept alongside the final score", row.base_score, 7.9)
check("sub-scores land in breakdown", row.breakdown,
      {"perfil_fit": 8, "location": 7, "entidad_fit": 9, "condiciones": 7})
check("flags land in flags", row.flags["catalan_imprescindible"], True)
check("extra output fields land in extra", row.extra["tipo_contrato"], "indefinido")
check("text[] columns are lists", isinstance(row.match_signals, list), True)
check("the run is recorded", row.run_id, "run7")
check("evaluated_at is preserved", row.evaluated_at, "2026-09-12T10:00:00")

# Every column named in ARCHITECTURE.md 5.3 must exist on the model.
EXPECTED_COLUMNS = {
    "user_id", "run_id", "job_key", "title", "company", "url", "source",
    "published", "location", "verdict", "score", "base_score", "breakdown",
    "one_liner", "match_signals", "gaps", "red_flags", "flags", "extra",
    "generated_text", "salary_range_market", "evaluated_at",
}
check("job_results has exactly the documented columns",
      set(JobResult.model_fields), EXPECTED_COLUMNS)
check("the row is JSON-serialisable for a jsonb column",
      isinstance(row.model_dump_json(), str), True)


# ===========================================================================
section("EXPORT — columns match the notebooks' own definitions")

for label, profile, cols, tracker in (("gabriel", gabriel, G_COLS, G_TRACKER),
                                      ("camila", camila, C_COLS, C_TRACKER)):
    main_sheet, tracker_sheet = profile.export.sheets
    check(f"{label}: main sheet labels match the notebook",
          [c.label for c in main_sheet.columns], [c[0] for c in cols])
    check(f"{label}: main sheet widths match the notebook",
          [c.width for c in main_sheet.columns], [c[1] for c in cols])
    check(f"{label}: tracker labels match the notebook",
          [c.label for c in tracker_sheet.columns], [c[0] for c in tracker])
    check(f"{label}: tracker widths match the notebook",
          [c.width for c in tracker_sheet.columns], [c[1] for c in tracker])
    check(f"{label}: URL is the last column of the main sheet",
          main_sheet.columns[-1].value, "field:url")


# ===========================================================================
section("EXPORT — rebuilt from the real rows of both .xlsx files")


def rows_from_xlsx(path: Path, profile, mapping: dict[str, str],
                   score_col: str, verdict_col: str, user_id: str) -> list[JobResult]:
    """Rebuild JobResult rows from an existing result file."""
    import re
    back = {v: k for k, v in {"YES": profile.verdict_labels.yes,
                              "MAYBE": profile.verdict_labels.maybe,
                              "NO": profile.verdict_labels.no}.items()}
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["Job Scout"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    out = []
    for index, raw in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        d = dict(zip(header, raw))
        breakdown = str(d.get(mapping["breakdown"]) or "")
        pairs = re.findall(r"(\w+):(\d+)", breakdown)
        try:
            score = float(d.get(score_col))
        except (TypeError, ValueError):
            continue
        label_to_key = {dim.label: dim.key for dim in profile.dimensions}
        out.append(JobResult(
            user_id=user_id, run_id="imported", job_key=f"k{index}",
            title=str(d.get(mapping["title"]) or ""),
            company=str(d.get(mapping["company"]) or ""),
            url=str(d.get(mapping["url"]) or f"https://imported/{index}"),
            source=str(d.get(mapping["source"]) or ""),
            verdict=back.get(str(d.get(verdict_col)), "NO"),
            score=score, base_score=score,
            breakdown={label_to_key[k]: int(v) for k, v in pairs if k in label_to_key},
            one_liner=str(d.get(mapping["one_liner"]) or ""),
        ))
    wb.close()
    return out


g_rows = rows_from_xlsx(
    REPO / "job_scout_results.xlsx", gabriel,
    {"breakdown": "Score breakdown", "title": "Titre", "company": "Entreprise",
     "url": "URL", "source": "Source", "one_liner": "One-liner"},
    "Score", "Verdict", "gabriel")
c_rows = rows_from_xlsx(
    REPO / "job_scout_camila" / "job_scout_camila_results.xlsx", camila,
    {"breakdown": "Desglose puntuación", "title": "Título", "company": "Entidad",
     "url": "Enlace", "source": "Fuente", "one_liner": "Resumen"},
    "Puntuación", "Veredicto", "camila")

print(f"\n  rebuilt {len(g_rows)} gabriel rows and {len(c_rows)} camila rows")

for label, profile, rows, cols, tracker in (
    ("gabriel", gabriel, g_rows, G_COLS, G_TRACKER),
    ("camila", camila, c_rows, C_COLS, C_TRACKER),
):
    user_store = JsonStore(TMP / "export")
    inserted = user_store.save_results(label, rows)
    stored = user_store.get_results(label)
    path = TMP / f"{label}.xlsx"
    written = export_xlsx(profile, stored, path)

    wb = load_workbook(path)
    ws = wb["Job Scout"]
    print(f"\n  {label}: {inserted} stored, wrote {written}")

    check(f"{label}: sheets are main + tracker + summary",
          wb.sheetnames, [profile.export.sheets[0].name,
                          profile.export.sheets[1].name, "Summary"])
    check(f"{label}: exported headers equal the notebook's labels",
          [c.value for c in ws[1]], [c[0] for c in cols])
    check(f"{label}: exported widths equal the notebook's widths",
          [ws.column_dimensions[chr(64 + i)].width for i in range(1, len(cols) + 1)],
          [float(c[1]) for c in cols])
    check(f"{label}: one row per stored result", ws.max_row - 1, len(stored))
    check(f"{label}: top row frozen", ws.freeze_panes, "A2")
    check(f"{label}: autofilter set", bool(ws.auto_filter.ref), True)

    # The verdict cell must be coloured from the INTERNAL verdict while
    # displaying the profile's own label.
    verdict_col = 1 + [c.value for c in profile.export.sheets[0].columns].index(
        "computed:verdict_display")
    displayed = {ws.cell(row=r, column=verdict_col).value for r in range(2, ws.max_row + 1)}
    expected_labels = {profile.verdict_labels.yes, profile.verdict_labels.maybe,
                       profile.verdict_labels.no}
    check(f"{label}: verdict column shows only this profile's labels",
          displayed - expected_labels, set())

    yes_rows = [r for r in range(2, ws.max_row + 1)
                if ws.cell(row=r, column=verdict_col).value == profile.verdict_labels.yes]
    if yes_rows:
        cell = ws.cell(row=yes_rows[0], column=verdict_col)
        check(f"{label}: a YES cell is green", cell.fill.start_color.rgb[-6:], "C6EFCE")
        check(f"{label}: and bold", cell.font.bold, True)

    tracker_ws = wb[profile.export.sheets[1].name]
    check(f"{label}: tracker headers equal the notebook's",
          [c.value for c in tracker_ws[1]], [c[0] for c in tracker])
    check(f"{label}: tracker holds only YES rows",
          tracker_ws.max_row - 1,
          len([r for r in stored if r.verdict == "YES"]))

    summary = wb["Summary"]
    check(f"{label}: summary total matches", summary["B3"].value, len(stored))
    check(f"{label}: summary counts the internal verdicts",
          [summary.cell(row=r, column=2).value for r in (4, 5, 6)],
          [sum(1 for x in stored if x.verdict == v) for v in ("YES", "MAYBE", "NO")])
    wb.close()


# ===========================================================================
section("EXPORT — computed values")

sample = JobResult(user_id="u", run_id="r", job_key="k", title="Junior AI Engineer",
                   url="https://x", verdict="YES", score=8.7, base_score=8.7,
                   published="2026-09-01",
                   breakdown={"tech_fit": 9, "location": 10,
                              "company_size": 8, "seniority": 9},
                   flags={"catalan_imprescindible": True})

check("score_breakdown uses the profile's short labels",
      resolve("computed:score_breakdown", gabriel, sample, None, "ts"),
      "Tech:9 Loc:10 Size:8 Sen:9 → 8.7")
check("verdict_display uses the profile's labels",
      resolve("computed:verdict_display", gabriel, sample, None, "ts"), "OUI")
check("and camila's are different",
      resolve("computed:verdict_display", camila,
              sample.model_copy(update={"verdict": "MAYBE"}), None, "ts"),
      "OPORTUNISTA")
check("priority uses gabriel's emoji",
      resolve("computed:priority", gabriel, sample, None, "ts"), "🔴")
check("priority uses camila's words",
      resolve("computed:priority", camila, sample, None, "ts"), "Alta")
check("default application status is per profile",
      resolve("computed:application_status", gabriel, sample, None, "ts"), "À postuler")
check("and camila's differs",
      resolve("computed:application_status", camila, sample, None, "ts"), "Pendiente")
check("an existing application overrides the default",
      resolve("computed:application_status", gabriel, sample,
              Application(user_id="u", job_url="https://x", status="Entretien"), "ts"),
      "Entretien")

check("target_salary picks the junior range for a junior title",
      resolve("computed:target_salary", gabriel, sample, None, "ts"),
      gabriel.target_salary_rule.junior_range)
check("and the default range otherwise",
      resolve("computed:target_salary", gabriel,
              sample.model_copy(update={"title": "Staff Engineer",
                                        "breakdown": {"seniority": 3}}), None, "ts"),
      gabriel.target_salary_rule.default_range)
check("camila has no salary rule, so the column is empty",
      resolve("computed:target_salary", camila, sample, None, "ts"), "")

check("flag: renders in the profile's language",
      resolve("flag:catalan_imprescindible", camila, sample, None, "ts"), "Sí")
check("and false flags too",
      resolve("flag:horario_atipico", camila, sample, None, "ts"), "No")
check("field: joins list columns",
      resolve("field:match_signals", gabriel,
              sample.model_copy(update={"match_signals": ["a", "b"]}), None, "ts"),
      "a | b")

check("priority thresholds: 8.5+ is high", priority_key(8.6, ""), "high")
check("7.5-8.4 is medium", priority_key(7.9, ""), "medium")
check("below 7.5 is low", priority_key(6.0, ""), "low")
check("a fresh posting is high regardless of score",
      priority_key(3.0, __import__("datetime").datetime.now().isoformat()), "high")

unknown = False
try:
    resolve("computed:nonexistent", gabriel, sample, None, "ts")
except KeyError:
    unknown = True
check("an unknown computed name raises", unknown, True)
check("every computed name used by both profiles is registered",
      sorted({c.value.split(":", 1)[1]
              for p in (gabriel, camila) for s in p.export.sheets
              for c in s.columns if c.value.startswith("computed:")}
             - set(REGISTRY)), [])


# ===========================================================================
print()
print("=" * 78)
if FAILURES:
    print(f"FAILED — {len(FAILURES)} of {CHECKS} checks:")
    for f in FAILURES:
        print(f"  - {f}")
    print("=" * 78)
    sys.exit(1)
print(f"PASSED — {CHECKS} checks. No API call, no network.")
print("=" * 78)
