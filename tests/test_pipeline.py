"""Block 6: the whole pipeline, exercised end to end without spending anything.

    python3 tests/test_pipeline.py

Costs nothing. The Anthropic client is replaced by a fake that returns canned
JSON, so every path through scoring and text generation runs for real except
the HTTP call itself. Job postings come from the saved connector dumps.

The point of the fake is not to mock away the interesting part: parsing,
verdict computation, retries, parallelism, persistence and export all execute.
Only the model's answer is substituted.
"""

from __future__ import annotations

import json
import shutil
import sys
import threading
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

from jobscout.connectors import Secrets  # noqa: E402
from jobscout.pipeline import (  # noqa: E402
    AWAITING_SELECTION, RunOptions, run,
)
from jobscout.profile import load_profile  # noqa: E402
from jobscout.store.base import Application  # noqa: E402
from jobscout.store.json_store import JsonStore  # noqa: E402
from jobscout.writing import build_writing_prompt, build_writing_system  # noqa: E402

from cost_guard import CostDecision, CostPolicy  # noqa: E402

CHECKS = 0
FAILURES: list[str] = []
TMP = Path("/private/tmp/claude-501/-Users-gabrielernoult-Desktop-GIT-Repos-Job-Search"
           "/a697afae-6c9e-4187-870d-45749f73a1a8/scratchpad/block6")


def check(label: str, actual, expected) -> None:
    global CHECKS
    CHECKS += 1
    if actual == expected:
        print(f"  OK    {label}")
        return
    FAILURES.append(label)
    print(f"  FAIL  {label}")
    print(f"          expected: {str(expected)[:200]!r}")
    print(f"          actual  : {str(actual)[:200]!r}")


def section(title: str) -> None:
    print(f"\n{'=' * 78}\n{title}\n{'=' * 78}")


shutil.rmtree(TMP, ignore_errors=True)
TMP.mkdir(parents=True, exist_ok=True)

gabriel = load_profile(REPO / "profiles" / "gabriel.json", base_dir=REPO)
camila = load_profile(REPO / "profiles" / "camila.json", base_dir=REPO)
DUMPS = str(REPO / "dumps")


# ---------------------------------------------------------------------------
# A fake Anthropic client
# ---------------------------------------------------------------------------

class FakeMessage:
    def __init__(self, text: str):
        self.content = [type("Block", (), {"text": text})()]


class FakeClient:
    """Returns canned answers and records every call.

    `scores` cycles through sub-score sets so a run produces a spread of
    verdicts rather than one value repeated.
    """

    def __init__(self, profile, scores=None, fail_first_n_json=0):
        self.profile = profile
        self.scores = scores or [
            {d.key: 9 for d in profile.dimensions},   # -> YES
            {d.key: 6 for d in profile.dimensions},   # -> MAYBE
            {d.key: 2 for d in profile.dimensions},   # -> NO
        ]
        self.fail_first_n_json = fail_first_n_json
        self.calls: list[dict] = []
        self.scoring_calls = 0
        self.writing_calls = 0
        self._lock = threading.Lock()
        self.messages = self

    def create(self, *, model, max_tokens, system, messages, **kwargs):
        with self._lock:
            self.calls.append({"model": model, "system": system,
                               "prompt": messages[0]["content"]})
            index = len(self.calls) - 1
            is_scoring = '"sub_scores"' in messages[0]["content"]
            if is_scoring:
                self.scoring_calls += 1
                n = self.scoring_calls
            else:
                self.writing_calls += 1

        if not is_scoring:
            return FakeMessage("Dear Hiring Team,\n\nGenerated body text.\n\n"
                               "I would love to discuss the role in more detail.")

        if n <= self.fail_first_n_json:
            return FakeMessage("{ this is not valid json")

        sub = self.scores[(n - 1) % len(self.scores)]
        payload = {
            "sub_scores": sub,
            "flags": {f.key: False for f in self.profile.llm_flags},
            "one_liner": "Canned assessment.",
            "match_signals": ["signal"],
            "gaps": ["gap"],
            "red_flags": [],
            "salary_range_market": "€40k-50k",
        }
        for field in self.profile.extra_output_fields:
            payload[field] = "value"
        return FakeMessage(json.dumps(payload))


def free_profile(profile):
    """Same profile with every Apify source replaying its dump: cost $0.00."""
    return profile.model_copy(update={"sources": [
        s.model_copy(update={"reuse_dump": True})
        if s.type in ("infojobs_apify", "linkedin_apify") else s
        for s in profile.sources
    ]})


# ===========================================================================
section("COST GUARD — the pipeline stops before anything is fetched or scored")

store = JsonStore(TMP / "held")
client = FakeClient(camila)
report = run(camila, store, Secrets(), RunOptions(dumps_dir=DUMPS),
             client=client, repo_dir=REPO)

check("status is the guard's decision", report.status,
      CostDecision.NEEDS_CONFIRMATION.value)
check("not ran", report.ran, False)
check("no job was scored", client.scoring_calls, 0)
check("no text was generated", client.writing_calls, 0)
check("no results stored", store.get_results("camila"), [])
check("the estimate is returned for the frontend",
      round(report.cost["usd_total"], 2), 1.38)
check("with a fingerprint to confirm", bool(report.cost["fingerprint"]), True)
check("the held run is recorded so confirmation can arrive later",
      store.get_run("camila", report.run_id).cost_fingerprint,
      report.cost["fingerprint"])

hard = camila.model_copy(update={
    "cost_policy": CostPolicy(max_cost_per_run=0.50, confirm_above=0.10)})
rejected = run(hard, JsonStore(TMP / "rejected"), Secrets(),
               RunOptions(dumps_dir=DUMPS), client=FakeClient(camila), repo_dir=REPO)
check("a run over the hard cap is rejected", rejected.status,
      CostDecision.REJECTED_OVER_HARD_CAP.value)


# ===========================================================================
section("FULL RUN — camila, every source replaying a dump (cost $0.00)")

store = JsonStore(TMP / "camila")
client = FakeClient(camila)
report = run(
    free_profile(camila), store, Secrets(),
    RunOptions(dumps_dir=DUMPS, max_jobs_to_score=30,
               export_xlsx_path=str(TMP / "camila.xlsx")),
    client=client, repo_dir=REPO,
)

print(f"\n  status: {report.status}")
print(f"  counts: {report.counts}")

check("the guard approved a free run", report.status, CostDecision.OK.value)
check("estimated cost is zero", report.cost["usd_total"], 0.0)
check("scoring ran for the capped number", client.scoring_calls, 30)
check("counts report what was scored", report.counts["scored"], 30)
check("and what the cap excluded", report.counts["capped_out"] > 0, True)
check("results were persisted", report.counts["saved"], 30)
check("the report carries the rows", len(report.results), 30)
check("stored rows match the report", len(store.get_results("camila")), 30)

verdicts = {r.verdict for r in report.results}
check("the run produced a spread of verdicts", sorted(verdicts),
      ["MAYBE", "NO", "YES"])
check("verdict counts add up",
      sum(report.counts[f"verdict_{v}"] for v in ("YES", "MAYBE", "NO")), 30)

row = report.results[0]
check("rows carry the run id", row.run_id, report.run_id)
check("rows carry the user", row.user_id, "camila")
check("breakdown is filled", len(row.breakdown), len(camila.dimensions))
check("job_key is set", len(row.job_key), 12)
check("extra output fields survive", row.extra.get("tipo_contrato"), "value")

check("seen cache was written", len(store.seen_keys("camila")), 30)
check("the run was recorded", store.get_run("camila", report.run_id).status, "OK")
check("with its counts", store.get_run("camila", report.run_id).counts["scored"], 30)

section("FULL RUN — text generation")

generated = [r for r in report.results if r.generated_text]
eligible = [r for r in report.results
            if r.verdict in camila.writing.generate_for_verdicts]
check("text generated for every eligible verdict",
      len(generated), len(eligible))
check("and only for those",
      {r.verdict for r in generated} - set(camila.writing.generate_for_verdicts),
      set())
check("the writing model was called that many times",
      client.writing_calls, len(eligible))
check("counts report it", report.counts["texts_generated"], len(eligible))
print(f"  {len(generated)} texts for verdicts {camila.writing.generate_for_verdicts}")

section("FULL RUN — export is optional and secondary")

check("the export path is reported", report.export_path.endswith("camila.xlsx"), True)
check("the file exists", Path(report.export_path).exists(), True)
from openpyxl import load_workbook  # noqa: E402
wb = load_workbook(report.export_path)
check("sheets as configured", wb.sheetnames, ["Job Scout", "Tracker", "Summary"])
exported = wb["Job Scout"].max_row - 1
check("only the configured verdicts are exported", exported,
      len([r for r in store.get_results("camila") if r.verdict in ("YES", "MAYBE")]))
wb.close()

no_export = run(free_profile(camila), JsonStore(TMP / "noexport"), Secrets(),
                RunOptions(dumps_dir=DUMPS, max_jobs_to_score=3),
                client=FakeClient(camila), repo_dir=REPO)
check("no export path means no file", no_export.export_path, "")
check("but job_results is still written", no_export.counts["saved"], 3)


# ===========================================================================
section("SECOND RUN — the seen cache and the unique URL constraint")

client2 = FakeClient(camila)
again = run(free_profile(camila), store, Secrets(),
            RunOptions(dumps_dir=DUMPS, max_jobs_to_score=30),
            client=client2, repo_dir=REPO)
print(f"  counts: {again.counts}")
check("previously seen jobs are not re-scored", client2.scoring_calls,
      again.counts["scored"])
check("the seen cache reports the skips", again.counts["already_seen"] > 0, True)
check("nothing already stored is saved twice",
      len(store.get_results("camila")),
      30 + again.counts["saved"])

section("SECOND RUN — an application survives it")

target = store.get_results("camila")[0]
store.save_application(Application(user_id="camila", job_url=target.url,
                                   status="Entrevista", notes="lunes 10h"))
run(free_profile(camila), store, Secrets(),
    RunOptions(dumps_dir=DUMPS, max_jobs_to_score=2),
    client=FakeClient(camila), repo_dir=REPO)
check("status untouched by a later run",
      store.get_applications("camila")[target.url].status, "Entrevista")
check("notes untouched",
      store.get_applications("camila")[target.url].notes, "lunes 10h")


# ===========================================================================
section("ROBUSTNESS — bad model output, and the retry")

# One bad answer, then a good one: the retry must rescue the job. Serial
# workers so the call order is deterministic.
retry_client = FakeClient(camila, fail_first_n_json=1)
retry_report = run(free_profile(camila), JsonStore(TMP / "retry"), Secrets(),
                   RunOptions(dumps_dir=DUMPS, max_jobs_to_score=5,
                              generate_text=False, scoring_workers=1),
                   client=retry_client, repo_dir=REPO)
print(f"  1 invalid answer: {retry_client.scoring_calls} calls for 5 jobs")
check("the retry costs one extra call", retry_client.scoring_calls, 6)
check("and no job is lost", retry_report.counts["scored"], 5)

# Both attempts bad: the job is dropped, the run survives, and the count says
# so rather than the job silently becoming a made-up middling score.
doomed_client = FakeClient(camila, fail_first_n_json=2)
doomed = run(free_profile(camila), JsonStore(TMP / "doomed"), Secrets(),
             RunOptions(dumps_dir=DUMPS, max_jobs_to_score=5,
                        generate_text=False, scoring_workers=1),
             client=doomed_client, repo_dir=REPO)
print(f"  2 invalid answers: {doomed.counts['scored']} scored, "
      f"{doomed.counts['unscoreable']} unscoreable")
check("a job whose retry also fails is dropped", doomed.counts["scored"], 4)
check("and is reported as unscoreable, not scored at a default",
      doomed.counts["unscoreable"], 1)
check("the run still completes", doomed.status, "OK")
check("and only the scored rows are persisted", doomed.counts["saved"], 4)

section("ROBUSTNESS — no global state between two profiles")

g_store = JsonStore(TMP / "shared")
c_store = JsonStore(TMP / "shared")
g_client = FakeClient(gabriel)
c_client = FakeClient(camila)
g_report = run(free_profile(gabriel), g_store, Secrets(),
               RunOptions(dumps_dir=DUMPS, max_jobs_to_score=4,
                          generate_text=False),
               client=g_client, repo_dir=REPO)
c_report = run(free_profile(camila), c_store, Secrets(),
               RunOptions(dumps_dir=DUMPS, max_jobs_to_score=4,
                          generate_text=False),
               client=c_client, repo_dir=REPO)

check("both runs succeeded",
      [g_report.status, c_report.status], ["OK", "OK"])
check("each scored its own jobs", [g_report.counts["scored"],
                                   c_report.counts["scored"]], [4, 4])
check("run ids differ", g_report.run_id != c_report.run_id, True)
check("gabriel's store holds only his rows",
      {r.user_id for r in g_store.get_results("gabriel")}, {"gabriel"})
check("camila's likewise",
      {r.user_id for r in c_store.get_results("camila")}, {"camila"})
check("gabriel's dimensions were used",
      sorted(g_report.results[0].breakdown), sorted(gabriel.dimension_keys))
check("camila's dimensions were used",
      sorted(c_report.results[0].breakdown), sorted(camila.dimension_keys))
check("each prompt carried its own candidate summary",
      [gabriel.candidate_summary[:30] in g_client.calls[0]["prompt"],
       camila.candidate_summary[:30] in c_client.calls[0]["prompt"]],
      [True, True])


# ===========================================================================
section("WRITING — one mechanism, two presets")

g_system = build_writing_system(gabriel)
c_system = build_writing_system(camila)

check("gabriel's length floor is stated", "AT LEAST 550 words" in g_system, True)
check("camila's ceiling is stated", "AT MOST 150 words" in c_system, True)
check("gabriel's mandatory closing sentence is verbatim",
      gabriel.writing.closing_sentence in g_system, True)
check("gabriel's verbatim Amazon phrase is required",
      gabriel.writing.verbatim_phrases[0] in g_system, True)
check("the 'One small aside' paragraph is required",
      "One small aside" in g_system, True)
check("em dashes are banned for gabriel", "em dashes" in g_system, True)
check("GitHub links are banned for gabriel", "GitHub links" in g_system, True)
check("camila's salutation template is present",
      "Hola equipo de" in c_system, True)
check("languages differ", ["English" in g_system, "Spanish" in c_system],
      [True, True])

scored_sample = report.results[0]
from jobscout.scoring import ScoredJob  # noqa: E402
sample = ScoredJob(title="Educador social", company="Fundació Alfa",
                   source="xarxanet", verdict="YES", one_liner="Buen encaje.",
                   match_signals=["comunitario"])
c_prompt = build_writing_prompt(camila, sample, "CV TEXT HERE")
check("the salutation is interpolated with the real company",
      "Hola equipo de Fundació Alfa," in c_prompt, True)
check("the CV is included", "CV TEXT HERE" in c_prompt, True)
check("camila does not fetch company context",
      camila.writing.include_company_context, False)
check("gabriel does", gabriel.writing.include_company_context, True)


# ===========================================================================
section("AVAILABLE_TO_SCORE — the real pool, named before truncation")

# gabriel replaying his saved LinkedIn dump: no Apify call, no Apify cost.
# RSS is live and free, so the absolute numbers move day to day; every check
# below is a relation between counts, never a hardcoded total.
replay = gabriel.model_copy(update={"sources": [
    src.model_copy(update={"reuse_dump": True})
    if src.type in ("linkedin_apify", "infojobs_apify") else src
    for src in gabriel.sources
]})

store = JsonStore(TMP / "pool")
client = FakeClient(gabriel)
report = run(replay, store, Secrets.from_env(),
             RunOptions(max_jobs_to_score=5, dumps_dir=DUMPS),
             user_id="pool-user", client=client, repo_dir=REPO)

counts = report.counts
pool = counts["available_to_score"]
check("the pool is reported", pool > 0, True)
check("and it is the number BEFORE the cap", pool > counts["to_score"], True)
check("the cap is what to_score reflects", counts["to_score"], 5)
check("pool = to_score + capped_out, the old implicit form",
      counts["to_score"] + counts["capped_out"], pool)
check("only the capped number was billed", client.scoring_calls, 5)

# Without a cap the pool and to_score agree, which is the other half of the
# contract the selector depends on.
store2 = JsonStore(TMP / "pool_uncapped")
client2 = FakeClient(gabriel)
report2 = run(replay, store2, Secrets.from_env(),
              RunOptions(max_jobs_to_score=None, dumps_dir=DUMPS,
                         generate_text=False),
              user_id="pool-user-2", client=client2, repo_dir=REPO)
check("with no cap, pool == to_score",
      report2.counts["available_to_score"], report2.counts["to_score"])
check("and nothing was capped out", "capped_out" in report2.counts, False)


section("AWAITING_SELECTION — collected and paid for, not yet scored")

store3 = JsonStore(TMP / "await")
client3 = FakeClient(gabriel)
held = run(replay, store3, Secrets.from_env(),
           RunOptions(select_after_collect=True, dumps_dir=DUMPS),
           user_id="await-user", client=client3, repo_dir=REPO)

check("the run stops at AWAITING_SELECTION", held.status, AWAITING_SELECTION)
check("NOT a cost-guard decision: the guard said OK",
      held.cost["decision"], CostDecision.OK.value)
check("no scoring call was made", client3.scoring_calls, 0)
check("no writing call either", client3.writing_calls, 0)
check("the pool is in the report", held.counts["available_to_score"] > 0, True)
check("and it matches an uncapped run of the same profile",
      held.counts["available_to_score"], report2.counts["available_to_score"])
check("the reason tells the caller what to do",
      "Choose how many to score" in held.reason, True)

record = store3.get_run("await-user", held.run_id)
check("the state is persisted, not held in memory", record is not None, True)
check("with the same status", record.status, AWAITING_SELECTION)
check("and the pool readable from the store",
      record.counts["available_to_score"], held.counts["available_to_score"])
check("nothing was written to job_results",
      len(store3.get_results("await-user")), 0)
check("and nothing was marked as seen, so a resume re-collects the same set",
      len(store3.seen_keys("await-user")), 0)

# Resuming: same run_id, a chosen volume, still no Apify call.
client4 = FakeClient(gabriel)
resumed = run(replay, store3, Secrets.from_env(),
              RunOptions(max_jobs_to_score=4, dumps_dir=DUMPS,
                         generate_text=False),
              user_id="await-user", client=client4, run_id=held.run_id,
              repo_dir=REPO)
check("the resumed run finishes", resumed.status, CostDecision.OK.value)
check("under the same run_id", resumed.run_id, held.run_id)
check("scoring exactly the chosen number", client4.scoring_calls, 4)
check("and the Apify cost is still zero", resumed.cost["usd_total"], 0.0)
check("results are saved this time", len(store3.get_results("await-user")) > 0, True)
check("the run record now reflects the finished run",
      store3.get_run("await-user", held.run_id).status, CostDecision.OK.value)


# ===========================================================================
print()
print("=" * 78)
if FAILURES:
    print(f"FAILED — {len(FAILURES)} of {CHECKS} checks:")
    for f in FAILURES:
        print(f"  - {f}")
    print("=" * 78)
    sys.exit(1)
print(f"PASSED — {CHECKS} checks. No Anthropic call, no Apify call.")
print("=" * 78)
