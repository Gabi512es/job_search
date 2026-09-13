"""Excel export — a convenience, not the deliverable.

The product of a run is the `job_results` table, which the app reads directly
(ARCHITECTURE.md 8). This module turns those same rows into a spreadsheet when
a user wants an offline copy.

The styling skeleton is the notebooks', kept as-is because it was already
identical in both: styled header, frozen top row, autofilter, thin borders,
verdict-coloured cells, and one summary sheet. What was hardcoded - which
columns, in which order, with which labels - now comes from the profile.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from jobscout.export.computed import resolve
from jobscout.profile import UserProfile
from jobscout.store.base import Application, JobResult

HEADER_BG = "1F4E79"
HEADER_FONT = "FFFFFF"

# Keyed by the internal verdict, so the colour does not depend on the display
# label. In the Camila notebook these had to be kept in sync by hand with a
# four-line comment explaining the trap.
VERDICT_COLORS = {
    "YES": {"bg": "C6EFCE", "font": "276221"},
    "MAYBE": {"bg": "FFEB9C", "font": "9C5700"},
    "NO": {"bg": "FFC7CE", "font": "9C0006"},
}
PRIORITY_COLORS = {
    "high": {"bg": "FFC7CE", "font": "9C0006"},
    "medium": {"bg": "FFD966", "font": "7F6000"},
    "low": {"bg": "FFF2CC", "font": "7F6000"},
}

THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BASE_FONT = Font(name="Arial", size=9)


def _write_header(ws, columns) -> None:
    fill = PatternFill("solid", start_color=HEADER_BG)
    font = Font(bold=True, color=HEADER_FONT, name="Arial", size=10)
    for index, column in enumerate(columns, 1):
        cell = ws.cell(row=1, column=index, value=column.label)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = column.width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"


def _write_sheet(
    wb: Workbook,
    sheet,
    profile: UserProfile,
    results: list[JobResult],
    applications: dict[str, Application],
    run_timestamp: str,
) -> int:
    rows = results
    if sheet.only_verdicts:
        rows = [r for r in rows if r.verdict in sheet.only_verdicts]

    ws = wb.create_sheet(sheet.name)
    _write_header(ws, sheet.columns)

    for row_index, result in enumerate(rows, 2):
        application = applications.get(result.url)
        colour = VERDICT_COLORS.get(result.verdict, {"bg": "FFFFFF", "font": "000000"})

        for col_index, column in enumerate(sheet.columns, 1):
            value = resolve(column.value, profile, result, application, run_timestamp)
            cell = ws.cell(row=row_index, column=col_index, value=value)
            cell.border = BORDER
            cell.font = BASE_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)

            if column.value == "computed:verdict_display":
                cell.fill = PatternFill("solid", start_color=colour["bg"])
                cell.font = Font(color=colour["font"], name="Arial", size=9,
                                 bold=(result.verdict == "YES"))
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif column.value == "computed:priority":
                from jobscout.export.computed import priority_key
                key = (application.priority if application and application.priority
                       else priority_key(result.score, result.published))
                swatch = PRIORITY_COLORS.get(key)
                if swatch:
                    cell.fill = PatternFill("solid", start_color=swatch["bg"])
                    cell.font = Font(color=swatch["font"], name="Arial", size=9, bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif column.value in ("computed:run_date", "field:score"):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif column.value == "field:url":
                cell.alignment = Alignment(vertical="top", wrap_text=False)

        ws.row_dimensions[row_index].height = 100 if result.generated_text else 45

    return len(rows)


def _write_summary(wb: Workbook, profile: UserProfile, results: list[JobResult]) -> None:
    """Counts come from the in-memory results, keyed on the INTERNAL verdict.

    The notebook counted by re-reading the display column, which meant the
    counters broke the moment the labels were translated.
    """
    ws = wb.create_sheet("Summary")
    ws["A1"] = "Last updated"
    ws["B1"] = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws["A3"] = "Total"
    ws["B3"] = len(results)
    for row, verdict in enumerate(("YES", "MAYBE", "NO"), 4):
        ws.cell(row=row, column=1, value=profile.verdict_labels.label(verdict))
        ws.cell(row=row, column=2, value=sum(1 for r in results if r.verdict == verdict))
    for column in ("A", "B"):
        ws.column_dimensions[column].width = 22


def export_xlsx(
    profile: UserProfile,
    results: list[JobResult],
    path: Path | str,
    applications: dict[str, Application] | None = None,
    run_timestamp: str | None = None,
) -> dict[str, int]:
    """Write one workbook from stored results. Returns rows written per sheet.

    Always writes a fresh file rather than appending to an existing one: the
    source of truth is `job_results`, so a stale spreadsheet has nothing worth
    preserving, and appending was where the old code accumulated three
    incompatible column formats in one file.
    """
    applications = applications or {}
    run_timestamp = run_timestamp or datetime.now().strftime("%Y-%m-%d %H:%M")

    wb = Workbook()
    wb.remove(wb.active)  # drop the default sheet

    written: dict[str, int] = {}
    for sheet in profile.export.sheets:
        written[sheet.name] = _write_sheet(
            wb, sheet, profile, results, applications, run_timestamp
        )

    _write_summary(wb, profile, results)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return written
